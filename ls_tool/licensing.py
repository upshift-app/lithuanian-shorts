"""Reads (and can create) data/licensing.xlsx - the external licence sheet.

The site holds no licensing data, so this stays a spreadsheet the team keeps by
hand. A film is matched either by ``film_id`` (the WordPress post id, exact) or by
``film_title``. Only ``licence_signed`` is used as a hard constraint.
"""
from __future__ import annotations

from typing import List, Optional

from openpyxl import Workbook, load_workbook

from .config import LICENSING_XLSX

COLUMNS = [
    "film_id",
    "film_title",
    "licence_signed",
    "licence_until",
    "rights_holder",
    "notes",
]

_TRUE = {"yes", "y", "true", "1", "taip", "signed", "pasirasyta", "pasirašyta", "x"}
_FALSE = {"no", "n", "false", "0", "ne", "nesirasyta", "nepasirašyta", "-"}


def parse_bool(value) -> Optional[bool]:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    v = str(value).strip().lower()
    if v in _TRUE:
        return True
    if v in _FALSE:
        return False
    return None


def _str(value) -> Optional[str]:
    if value is None:
        return None
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    return text or None


def load_licensing(path=LICENSING_XLSX) -> List[dict]:
    if not path.exists():
        return []
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    index = {name: header.index(name) for name in COLUMNS if name in header}
    if "film_title" not in index and "film_id" not in index:
        raise ValueError(
            f"{path.name} needs at least a 'film_id' or 'film_title' column. "
            f"Found: {header}"
        )

    out = []
    for row in rows[1:]:
        if row is None or all(c is None or c == "" for c in row):
            continue

        def cell(name):
            i = index.get(name)
            return row[i] if i is not None and i < len(row) else None

        out.append({
            "film_id": cell("film_id"),
            "film_title": _str(cell("film_title")),
            "licence_signed": parse_bool(cell("licence_signed")),
            "licence_until": _str(cell("licence_until")),
            "rights_holder": _str(cell("rights_holder")),
            "notes": _str(cell("notes")),
        })
    return out


def write_template(catalog=None, path=LICENSING_XLSX, overwrite: bool = False):
    """Create the licence sheet, pre-filled with every film in the archive."""
    if path.exists() and not overwrite:
        raise FileExistsError(f"{path} already exists (pass overwrite=True to replace)")

    wb = Workbook()
    ws = wb.active
    ws.title = "Licences"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)

    if catalog is not None:
        for film in sorted(catalog.films, key=lambda f: (f.title or "").lower()):
            ws.append([film.id, film.title, "", "", "", ""])

    widths = {"A": 10, "B": 46, "C": 15, "D": 14, "E": 28, "F": 40}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
