"""Word (.docx) output - the editable format the team asked for.

Every document can be produced in Lithuanian or English (``lang="lt"|"en"``):
the chrome comes from ls_tool.i18n, the film content from the English
counterpart pages the scraper collects.

The layout follows the agency's chosen PDF design (see pdf_designs.py): logo
with the agency name under it, company details at the right edge, a navy rule,
and each film in a plain card with its still on the left.
"""
from __future__ import annotations

import datetime as _dt
import time
from pathlib import Path
from typing import List, Optional, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .config import ASSETS_DIR, OUTPUT_DIR
from .export_pdf import (STILL_BUDGET, _cropped, _film_meta, _headline,
                         _still_bytes)
from .i18n import (DEFAULT_LANG, company_lines, film_genre, film_synopsis,
                   film_title, minutes_label, norm_lang, t)
from .programme import Programme
from .screenings import FilmReport, totals

NAVY = RGBColor(0x1A, 0x1D, 0x45)
INK = RGBColor(0x14, 0x14, 0x14)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
NAVY_HEX, RULE_HEX = "1A1D45", "D8D8D8"
CARD_HEX, BAND_HEX, ZEBRA_HEX = "FAFAFA", "F2F2F2", "F7F7F7"
BODY_FONT = "Calibri"
LOGO = ASSETS_DIR / "logo.png"

MARGIN = Mm(22)
CONTENT = Mm(210 - 44)
STILL_W = Mm(50)
GAP = Mm(6)


# ------------------------------------------------------------------- low level

def _setup(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = MARGIN
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(22)
    _footer(sec)


def _footer(section) -> None:
    p = section.footer.paragraphs[0]
    p.text = ""
    r = p.add_run("Lithuanian Shorts\t")
    r.font.size, r.font.color.rgb = Pt(7.5), MUTED
    num = p.add_run()
    num.font.size, num.font.color.rgb = Pt(7.5), MUTED
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        num._r.append(el)


def _para(doc, text="", size=10.5, bold=False, italic=False, color=INK,
          space_after=4, align=None):
    p = doc.add_paragraph()
    _fill(p, text, size, bold, italic, color, space_after, align)
    return p


def _fill(p, text, size=10.5, bold=False, italic=False, color=INK,
          space_after=4, align=None):
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold, run.italic = bold, italic
    run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    return p


def _cell_para(cell, text, first=False, **kw):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    return _fill(p, text, **kw)


def _shade(cell, hex_color: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def _borders(parent, **edges) -> None:
    """edges: top/left/bottom/right/insideH/insideV -> (size_eighths, hex) or
    None for no line. ``parent`` is a cell or a table."""
    if hasattr(parent, "_tc"):
        pr, tag = parent._tc.get_or_add_tcPr(), "w:tcBorders"
    else:
        pr, tag = parent._tbl.tblPr, "w:tblBorders"
    box = OxmlElement(tag)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if name not in edges:
            continue
        el = OxmlElement(f"w:{name}")
        spec = edges[name]
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), spec[1])
        box.append(el)
    pr.append(box)


def _pad(cell, top=0, left=0, bottom=0, right=0) -> None:
    mar = OxmlElement("w:tcMar")
    for name, v in (("top", top), ("left", left), ("bottom", bottom),
                    ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(int(v)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    cell._tc.get_or_add_tcPr().append(mar)


def _table(doc, rows, widths):
    table = doc.add_table(rows=rows, cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
    _borders(table, top=None, left=None, bottom=None, right=None,
             insideH=None, insideV=None)
    return table


def _no_split(row) -> None:
    pr = row._tr.get_or_add_trPr()
    pr.append(OxmlElement("w:cantSplit"))


def _spacer(doc, pts: float) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("")
    r.font.size = Pt(max(pts, 1))
    p.paragraph_format.line_spacing = Pt(max(pts, 1))


def _fmt_eur(amount: float) -> str:
    return f"{amount:,.2f} EUR".replace(",", " ")


# ----------------------------------------------------------------------- header

def _header(doc, label: str, company: Optional[dict], lang: str) -> None:
    """Logo and agency name on the left, company details on the right, then
    the short navy rule."""
    lines = company_lines(company, lang, split_contact=True)
    table = _table(doc, 1, [CONTENT * 0.55, CONTENT * 0.45])
    left, right = table.rows[0].cells
    first = True
    if LOGO.is_file():
        left.paragraphs[0].add_run().add_picture(str(LOGO), width=Mm(54))
        left.paragraphs[0].paragraph_format.space_after = Pt(6)
        first = False
    _cell_para(left, label, first=first, size=9, bold=True, color=NAVY,
               space_after=0)
    for i, line in enumerate(lines):
        _cell_para(right, line, first=i == 0, size=8, color=MUTED,
                   space_after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _spacer(doc, 8)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(14)
    rule.paragraph_format.right_indent = CONTENT - Mm(22)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "18"), ("space", "1"),
                 ("color", NAVY_HEX)):
        bottom.set(qn(f"w:{k}"), v)
    pbdr.append(bottom)
    rule._p.get_or_add_pPr().append(pbdr)


def _strip(doc, pairs: List[tuple], shaded: bool, big: bool = False) -> None:
    """A row of label/value pairs: light rules above and below, or grey tiles."""
    table = _table(doc, 2, [CONTENT / len(pairs)] * len(pairs))
    for i, (label, value) in enumerate(pairs):
        top, bot = table.rows[0].cells[i], table.rows[1].cells[i]
        _cell_para(top, label.upper() if label else "", first=True, size=7.5,
                   bold=True, color=MUTED, space_after=0)
        _cell_para(bot, value, first=True, size=17 if big else 10, bold=True,
                   color=INK, space_after=0)
        for c, line in ((top, "top"), (bot, "bottom")):
            _pad(c, top=50 if c is top else 20, left=160 if shaded else 0,
                 bottom=50 if c is bot else 0)
            if shaded:
                _shade(c, BAND_HEX)
            else:
                _borders(c, **{line: (5, RULE_HEX)})


# --------------------------------------------------------------- programme doc

def _film_card(doc, n: int, film, lang: str, deadline: float) -> None:
    data = None
    url = getattr(film, "image", None)
    if url and str(url).startswith(("http://", "https://")) \
            and time.monotonic() <= deadline:
        data = _still_bytes(str(url))
    still = None
    if data:
        try:
            still = _cropped(data)
        except Exception:
            still = None

    widths = [STILL_W + GAP, CONTENT - STILL_W - GAP] if still else [CONTENT]
    table = _table(doc, 1, widths)
    _borders(table, top=(5, RULE_HEX), left=(5, RULE_HEX),
             bottom=(5, RULE_HEX), right=(5, RULE_HEX))
    _no_split(table.rows[0])
    for cell in table.rows[0].cells:
        _shade(cell, CARD_HEX)
    text_cell = table.rows[0].cells[-1]
    if still:
        img = table.rows[0].cells[0]
        _pad(img, top=170, left=200, bottom=170, right=0)
        img.paragraphs[0].add_run().add_picture(still, width=STILL_W)
        img.paragraphs[0].paragraph_format.space_after = Pt(0)
    _pad(text_cell, top=170, left=0 if still else 200, bottom=170, right=140)

    _cell_para(text_cell, _headline(film, lang), first=True, size=11.5,
               bold=True, space_after=2)
    meta = _film_meta(film, lang)
    if meta:
        _cell_para(text_cell, meta, size=8.8, color=MUTED, space_after=4)
    synopsis = film_synopsis(film, lang)
    if synopsis:
        _cell_para(text_cell, synopsis, size=9.6, space_after=0)


def programme_docx(proposals: Sequence[Programme], alternates=(),
                   path: Optional[Path] = None, subtitle: Optional[str] = None,
                   lang: Optional[str] = None,
                   company: Optional[dict] = None) -> Path:
    req = proposals[0].request if proposals else None
    lang = norm_lang(lang or (req.lang if req else DEFAULT_LANG))

    doc = Document()
    _setup(doc)
    heading = (req.title if req and req.title else None) or t(lang, "programme_title")
    deadline = time.monotonic() + STILL_BUDGET

    _header(doc, t(lang, "agency"), company, lang)

    for idx, prog in enumerate(proposals, 1):
        if idx > 1:
            doc.add_page_break()
        title = heading if len(proposals) == 1 \
            else f"{heading} - {t(lang, 'variant', n=idx)}"
        _para(doc, title, size=24, bold=True, space_after=6)

        intro = subtitle or (req.intro if req else None) or (req.occasion if req else None)
        if intro:
            for para in str(intro).splitlines():
                if para.strip():
                    _para(doc, para.strip(), size=10.5, space_after=4)
            _spacer(doc, 6)

        facts = [(t(lang, "duration_label").rstrip(":"),
                  f"{int(round(prog.total_minutes))} {t(lang, 'minute')}."),
                 (t(lang, "films_in_programme").rstrip(":"), str(len(prog.films)))]
        if req and req.rating:
            facts.append(("", req.rating))
        _strip(doc, facts, shaded=False)
        _spacer(doc, 14)

        for n, film in enumerate(prog.films, 1):
            if n > 1:
                _spacer(doc, 14)
            _film_card(doc, n, film, lang, deadline)

        warnings = prog.warnings(lang)
        if warnings:
            _spacer(doc, 14)
            _para(doc, t(lang, "notes").upper(), size=7.5, bold=True,
                  color=MUTED, space_after=2)
            for w in warnings:
                _para(doc, "• " + w, size=8.5, color=MUTED, space_after=1)

    if alternates:
        _spacer(doc, 14)
        _para(doc, t(lang, "alternates").upper(), size=7.5, bold=True,
              color=MUTED, space_after=2)
        _para(doc, t(lang, "alternates_note"), size=8.5, color=MUTED,
              space_after=3)
        for film in alternates:
            bits = " · ".join(str(x) for x in (
                film.year, film_genre(film, lang),
                minutes_label(film.duration_min, lang)) if x)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run("• " + film_title(film, lang))
            r.bold, r.font.size = True, Pt(8.5)
            r2 = p.add_run(" - " + bits)
            r2.font.size, r2.font.color.rgb = Pt(8.5), MUTED

    path = Path(path) if path else OUTPUT_DIR / _default_name(heading, "docx", lang)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


# ------------------------------------------------------------------- report doc

def report_docx(reports: Sequence[FilmReport], path: Optional[Path] = None,
                title: Optional[str] = None, period: Optional[str] = None,
                recipient: Optional[str] = None,
                lang: Optional[str] = None,
                company: Optional[dict] = None) -> Path:
    lang = norm_lang(lang)
    title = title or t(lang, "report_title")

    doc = Document()
    _setup(doc)

    _header(doc, t(lang, "agency"), company, lang)
    _para(doc, title, size=24, bold=True, space_after=4)
    if recipient:
        _para(doc, recipient, size=10.5, space_after=4)
    line = f"{t(lang, 'prepared')} {_dt.date.today():%Y-%m-%d}"
    if period:
        line += f" · {t(lang, 'period')} {period}"
    _para(doc, line, size=8.8, color=MUTED, space_after=10)

    agg = totals(reports)
    _strip(doc, [(t(lang, "th_films"), str(agg["films"])),
                 (t(lang, "th_events"), str(agg["events"])),
                 (t(lang, "th_screenings"), str(agg["screenings"])),
                 (t(lang, "th_revenue"), _fmt_eur(agg["revenue"]))],
           shaded=True, big=True)
    _spacer(doc, 14)

    for report in reports:
        _para(doc, report.display_title(lang), size=13.5, bold=True,
              space_after=2)
        meta = " · ".join(str(x) for x in (
            (f"{t(lang, 'dir_prefix')} {report.director}") if report.director else None,
            report.year,
            minutes_label(report.duration_min, lang) if report.duration_min else None,
        ) if x)
        if meta:
            _para(doc, meta, size=8.8, color=MUTED, space_after=3)
        facts = [f"{t(lang, 'n_screenings')}: {report.total_screenings}",
                 f"{t(lang, 'n_events')}: {report.total_events}",
                 f"{t(lang, 'th_revenue')}: {_fmt_eur(report.total_revenue)}"]
        if report.countries:
            facts.append(f"{t(lang, 'countries')}: " + ", ".join(report.countries))
        if report.period:
            facts.append(f"{t(lang, 'period_label')}: {report.period}")
        _para(doc, " · ".join(facts), size=9.6, space_after=6)

        widths = [CONTENT * w for w in (0.12, 0.30, 0.28, 0.15, 0.15)]
        n_rows = len(report.screenings) + 2
        table = _table(doc, n_rows, widths)
        _borders(table, top=(5, RULE_HEX), left=(5, RULE_HEX),
                 bottom=(5, RULE_HEX), right=(5, RULE_HEX))
        heads = (t(lang, "th_date"), t(lang, "th_event"), t(lang, "th_venue"),
                 t(lang, "th_screenings"), t(lang, "th_revenue"))
        body = [(s.date_label or "-", s.event or "-", s.place or "-",
                 str(s.screenings), _fmt_eur(s.revenue))
                for s in report.screenings]
        total = ("", "", t(lang, "total"), str(report.total_screenings),
                 _fmt_eur(report.total_revenue))
        for r, values in enumerate([heads] + body + [total]):
            row = table.rows[r]
            _no_split(row)
            bold = r == 0 or r == n_rows - 1
            for c, text in zip(row.cells, values):
                _pad(c, top=70, left=100, bottom=70, right=60)
                _cell_para(c, text, first=True, size=8.6, bold=bold,
                           space_after=0)
                if r == 0:
                    _borders(c, bottom=(8, "141414"))
                elif r == n_rows - 1:
                    _borders(c, top=(6, "141414"))
                elif r % 2 == 0:
                    _shade(c, ZEBRA_HEX)

        notes = [s.notes for s in report.screenings if s.notes]
        if notes:
            _para(doc, f"{t(lang, 'notes_prefix')}: " + "; ".join(notes),
                  size=8.5, italic=True, color=MUTED, space_after=4)
        _spacer(doc, 16)

    path = Path(path) if path else OUTPUT_DIR / _default_name(title, "docx", lang)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def _default_name(title: str, ext: str, lang: str = DEFAULT_LANG) -> str:
    import re
    from .catalog import fold
    slug = re.sub(r"[^a-z0-9]+", "-", fold(title)).strip("-") or "dokumentas"
    return f"{_dt.date.today():%Y%m%d}-{slug[:60]}-{lang}.{ext}"
